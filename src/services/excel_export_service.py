import re
import unicodedata
from datetime import date, datetime, time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def slugify(text: str) -> str:
    """Convierte un texto en un slug limpio para nombres de archivo."""
    if not text:
        return "sin-catedra"
    nfkd = unicodedata.normalize('NFKD', text).encode('ASCII', 'ignore').decode('utf-8')
    slug = re.sub(r'[^a-zA-Z0-9]+', '-', nfkd).strip('-').lower()
    return slug or "catedra"


def generar_nombre_sugerido(fecha_val, catedra: str, extension: str) -> str:
    """
    Genera el nombre sugerido oficial:
    asistencia_YYYY-MM-DD_<catedra-slug>.<ext>
    """
    if isinstance(fecha_val, (date, datetime)):
        fecha_str = fecha_val.strftime("%Y-%m-%d")
    else:
        fecha_str = str(fecha_val or date.today().strftime("%Y-%m-%d"))

    slug = slugify(catedra)
    ext = extension.lstrip('.')
    return f"asistencia_{fecha_str}_{slug}.{ext}"


class ExcelExportService:
    """
    Servicio de exportación a Excel (.xlsx) usando openpyxl.
    Cumple con el diseño, ordenamiento y casos especiales de la rúbrica del Incremento 6.
    """

    @staticmethod
    def exportar(ruta_archivo: str, detalle_jornada: tuple, asistencias: list):
        """
        Genera la planilla Excel de asistencia.

        detalle_jornada: (id, fecha, catedra, horario_entrada, horario_salida, nombre_usuario)
        asistencias: lista de tuplas (apellido, nombre, dni, libreta, entrada, salida)
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Planilla de Asistencia"

        # Asegurar líneas de cuadrícula visibles en Excel
        ws.views.sheetView[0].showGridLines = True

        # Desempaquetar datos de jornada
        _, fecha_val, catedra, h_entrada, h_salida, usuario_docente = detalle_jornada

        # Formateo de fechas y horas para el encabezado
        fecha_str = fecha_val.strftime("%d/%m/%Y") if isinstance(fecha_val, (date, datetime)) else str(fecha_val or "")
        
        def _fmt_hora(h):
            if isinstance(h, (time, datetime)):
                return h.strftime("%H:%M")
            return str(h or "")

        entrada_str = _fmt_hora(h_entrada)
        salida_str = _fmt_hora(h_salida)

        # ── Estilos ──────────────────────────────────────────────────
        fuente_titulo = Font(name="Calibri", size=16, bold=True, color="1F497D")
        fuente_meta_label = Font(name="Calibri", size=11, bold=True, color="333333")
        fuente_meta_val = Font(name="Calibri", size=11, color="111111")
        
        fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        fuente_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        fill_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        fill_blanco = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        
        borde_fino = Side(border_style="thin", color="D9D9D9")
        borde_celda = Border(left=borde_fino, right=borde_fino, top=borde_fino, bottom=borde_fino)
        
        borde_header = Border(
            left=Side(border_style="thin", color="1F497D"),
            right=Side(border_style="thin", color="1F497D"),
            top=Side(border_style="medium", color="1F497D"),
            bottom=Side(border_style="medium", color="1F497D")
        )

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        # ── 1. Encabezado de la Jornada ──────────────────────────────
        ws["B2"] = "PLANILLA DE ASISTENCIA"
        ws["B2"].font = fuente_titulo

        metadatos = [
            ("Fecha:", fecha_str),
            ("Cátedra:", str(catedra or "")),
            ("Horario de entrada:", entrada_str),
            ("Horario de salida:", salida_str),
            ("Usuario:", str(usuario_docente or ""))
        ]

        fila_actual = 4
        for label, val in metadatos:
            ws.cell(row=fila_actual, column=2, value=label).font = fuente_meta_label
            ws.cell(row=fila_actual, column=3, value=val).font = fuente_meta_val
            fila_actual += 1

        fila_actual += 1  # Espacio en blanco

        # ── 2. Tabla de Asistencia ───────────────────────────────────
        columnas_tabla = ["#", "Apellido", "Nombre", "DNI", "Libreta", "Entrada", "Salida"]
        
        for col_idx, col_name in enumerate(columnas_tabla, start=2):
            cell = ws.cell(row=fila_actual, column=col_idx, value=col_name)
            cell.font = fuente_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = borde_header

        fila_header = fila_actual
        fila_actual += 1

        total_alumnos = len(asistencias)

        if total_alumnos == 0:
            # Caso especial 7.2: Jornada sin asistencias
            ws.merge_cells(start_row=fila_actual, start_column=2, end_row=fila_actual, end_column=8)
            aviso_cell = ws.cell(row=fila_actual, column=2, value="No se registraron asistencias para esta Jornada.")
            aviso_cell.font = Font(name="Calibri", size=11, italic=True, color="7F7F7F")
            aviso_cell.alignment = align_center
            for c in range(2, 9):
                ws.cell(row=fila_actual, column=c).border = borde_celda
            fila_actual += 2
        else:
            # Filas de alumnos presentes
            for idx, alu in enumerate(asistencias, start=1):
                apellido, nombre, dni, libreta, f_ent, f_sal = alu

                ent_str = f_ent.strftime("%H:%M") if isinstance(f_ent, (datetime, time)) else (str(f_ent) if f_ent else "")
                sal_str = f_sal.strftime("%H:%M") if isinstance(f_sal, (datetime, time)) else ("Sin registrar" if f_sal is None else str(f_sal))

                valores_fila = [idx, apellido, nombre, dni, libreta, ent_str, sal_str]
                fill_fila = fill_zebra if idx % 2 == 0 else fill_blanco

                for col_idx, val in enumerate(valores_fila, start=2):
                    cell = ws.cell(row=fila_actual, column=col_idx, value=val)
                    cell.font = Font(name="Calibri", size=11)
                    cell.fill = fill_fila
                    cell.border = borde_celda
                    
                    # Alineación: centro para #, DNI, Libreta, Horas; izquierda para nombres
                    if col_idx in (2, 5, 6, 7, 8):
                        cell.alignment = align_center
                    else:
                        cell.alignment = align_left

                fila_actual += 1
            fila_actual += 1

        # ── 3. Contador Final de Alumnos ─────────────────────────────
        # Caso especial 7.3: Cuidar singular y plural
        if total_alumnos == 0:
            texto_contador = "0 alumnos presentes."
        elif total_alumnos == 1:
            texto_contador = "Se registró 1 alumno presente en total."
        else:
            texto_contador = f"Se registraron {total_alumnos} alumnos presentes en total."

        cell_contador = ws.cell(row=fila_actual, column=2, value=texto_contador)
        cell_contador.font = Font(name="Calibri", size=11, bold=True, color="1F497D")

        # ── 4. Ajuste automático de anchos de columna ────────────────
        anchos_minimos = {
            2: 6,   # #
            3: 18,  # Apellido
            4: 18,  # Nombre
            5: 14,  # DNI
            6: 14,  # Libreta
            7: 12,  # Entrada
            8: 14,  # Salida
        }

        for col_idx in range(2, 9):
            col_letter = get_column_letter(col_idx)
            max_len = anchos_minimos.get(col_idx, 12)
            
            # Revisar las celdas de la tabla para asegurar que no se corten
            for row in range(fila_header, fila_actual):
                val_str = str(ws.cell(row=row, column=col_idx).value or "")
                if len(val_str) > max_len and not ws.cell(row=row, column=col_idx).coordinate in ws.merged_cells:
                    max_len = len(val_str)
                    
            ws.column_dimensions[col_letter].width = max_len + 3

        # Margen de columna A
        ws.column_dimensions["A"].width = 4

        wb.save(ruta_archivo)
